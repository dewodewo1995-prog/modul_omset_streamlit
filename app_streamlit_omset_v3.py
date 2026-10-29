import calendar
import os
import shutil as sh
import tempfile
import zipfile
import openpyxl
import pandas as pd
import rarfile
import streamlit as st

# Konfigurasi halaman web Streamlit
st.set_page_config(
    page_title="Rekap Omset Web", page_icon="📊", layout="wide"
)

st.title("📊 Aplikasi Rekap Omset (Versi Web Cloud)")
st.write(
    "Upload file arsip (ZIP/RAR) atau file-file Excel harian langsung melalui"
    " browser!"
)


# --- TATA CARA CEPAT PENGGUNAAN (EXPANDER) ---
with st.expander("📖 Tata Cara Cepat Penggunaan (Klik untuk Membaca)", expanded=False):
  st.markdown("""
    1. **Format Penamaan File Harian**: Pastikan file Excel harian menggunakan format penamaan `tanggal_3hurufbulan` (Contoh: `1_okt.xlsx`, `2_okt.xlsx`, dst).
    2. **Langkah Kerja**: 
       * **Langkah 1**: Lakukan upload data harian terlebih dahulu (bisa pilih metode *Upload File Arsip ZIP/RAR* atau *Banyak File Excel*).
       * **Langkah 2**: Atur parameter (Jumlah Tanggal proses, Bulan, dan Tahun) pada bagian bawah setelah file ter-upload.
    3. **Menu Proses Baru**: Pilih tab *Proses Baru (Upload Web)* untuk membuat rekap omset dari awal menggunakan master template. Kolom tanggal (`DATE` mulai dari sel C8) akan otomatis ter-generate penuh dari tanggal 1 sampai akhir bulan (30/31 hari).
    4. **Menu Update Excel Existing**: Pilih tab *Update Excel Existing (Web)* untuk memperbarui file rekap omset yang sudah pernah dibuat sebelumnya.
    5. **Eksekusi Proses**: Klik tombol **Jalankan Proses Penggabungan** atau **Jalankan Update Excel Target** untuk memproses dan mendownload hasilnya.
    """)

st.markdown("---")


# --- FUNGSI PROSES DATA DARI DIREKTORI SEMENTARA ---
def proses_gabung_data_dari_folder(source_dir, jml_tgl, bln):
  df_list = []
  log_messages = []

  for i in range(1, jml_tgl + 1):
    file_base = os.path.join(source_dir, f"{i}_{bln}")
    file_path_xlsx = file_base + ".xlsx"
    file_path_xls = file_base + ".xls"

    target_file = None
    if os.path.exists(file_path_xlsx):
      target_file = file_path_xlsx
    elif os.path.exists(file_path_xls):
      target_file = file_path_xls

    if target_file:
      try:
        df = pd.read_excel(target_file, sheet_name="Sheet1", header=0)
        df["tanggal"] = i
        df_list.append(df)
        log_messages.append(f"✅ Berhasil memproses tanggal: {i}")
      except Exception as e:
        log_messages.append(f"❌ Error membaca file tanggal {i}: {e}")
    else:
      log_messages.append(f"❌ File Excel tanggal {i}_{bln} tidak ditemukan.")

  if not df_list:
    return None, log_messages

  df_gabungan = pd.concat(df_list, axis=0, ignore_index=True)
  return df_gabungan, log_messages


# --- AREA TAMPILAN UTAMA (TABS) ---
tab1, tab2 = st.tabs(
    ["🚀 Proses Baru (Upload Web)", "🔄 Update Excel Existing (Web)"]
)

# ================= TAB 1: PROSES BARU =================
with tab1:
  st.subheader("Proses Penggabungan & Template via Web")
  st.write(
      "**Langkah 1**: Upload data harian Anda di bawah ini, lalu tentukan"
      " parameter pengaturannya."
  )

  metode_upload = st.radio(
      "Pilih Metode Upload Data Harian (Tab 1):",
      ["📁 Upload File Arsip (ZIP / RAR)", "📄 Upload Banyak File Excel (.xlsx/.xls)"],
      key="metode_tab1",
  )

  uploaded_archive = None
  uploaded_files = None

  if metode_upload == "📁 Upload File Arsip (ZIP / RAR)":
    uploaded_archive = st.file_uploader(
        "Upload file .zip atau .rar yang berisi file-file harian:",
        type=["zip", "rar"],
        key="archive_tab1",
    )
  else:
    uploaded_files = st.file_uploader(
        "Upload file-file Excel harian (bisa pilih banyak sekaligus):",
        type=["xlsx", "xls"],
        accept_multiple_files=True,
        key="files_tab1",
    )

  st.markdown("")

  # --- PARAMETER PENGATURAN (DIBAWAH UPLOAD TAB 1) ---
  with st.container(border=True):
    st.markdown(
        "<h6 style='text-align: center; margin-bottom: 5px;'>⚙️ Parameter"
        " Pengaturan Data</h6>",
        unsafe_allow_html=True,
    )
    p1, p2, p3 = st.columns(3)

    with p1:
      jumlah_tanggal_t1 = st.selectbox(
          "Jumlah Tanggal (Proses):",
          options=list(range(1, 32)),
          index=0,
          key="jml_t1",
      )

    with p2:
      bulan_dict = {
          "Januari": ("jan", 1),
          "Februari": ("feb", 2),
          "Maret": ("mar", 3),
          "April": ("apr", 4),
          "Mei": ("mei", 5),
          "Juni": ("jun", 6),
          "Juli": ("jul", 7),
          "Agustus": ("agu", 8),
          "September": ("sep", 9),
          "Oktober": ("okt", 10),
          "November": ("nov", 11),
          "Desember": ("des", 12),
      }
      pilihan_bulan_nama_t1 = st.selectbox(
          "Bulan:", options=list(bulan_dict.keys()), key="bln_t1"
      )
      nama_bulan_t1, angka_bulan_t1 = bulan_dict[pilihan_bulan_nama_t1]

    with p3:
      pilihan_tahun_t1 = st.selectbox(
          "Tahun:",
          options=[2026, 2027, 2028, 2029, 2030],
          index=0,
          key="thn_t1",
      )

  st.markdown("")
  if st.button("🚀 Jalankan Proses Penggabungan", type="primary"):
    with tempfile.TemporaryDirectory() as temp_dir:
      sumber_data_siap = False

      if (
          metode_upload == "📁 Upload File Arsip (ZIP / RAR)"
          and uploaded_archive is not None
      ):
        with st.spinner("Mengekstrak file arsip..."):
          archive_path = os.path.join(temp_dir, uploaded_archive.name)
          with open(archive_path, "wb") as f:
            f.write(uploaded_archive.getbuffer())

          if uploaded_archive.name.endswith(".zip"):
            with zipfile.ZipFile(archive_path, "r") as zip_ref:
              zip_ref.extractall(temp_dir)
            sumber_data_siap = True
          elif uploaded_archive.name.endswith(".rar"):
            try:
              with rarfile.RarFile(archive_path, "r") as rar_ref:
                rar_ref.extractall(temp_dir)
              sumber_data_siap = True
            except Exception as e:
              st.error(
                  f"❌ Gagal mengekstrak file RAR. Pastikan sistem mendukung unrar: {e}"
              )

      elif (
          metode_upload == "📄 Upload Banyak File Excel (.xlsx/.xls)"
          and uploaded_files
      ):
        with st.spinner("Menyiapkan file yang di-upload..."):
          for uf in uploaded_files:
            temp_file_path = os.path.join(temp_dir, uf.name)
            with open(temp_file_path, "wb") as f:
              f.write(uf.getbuffer())
          sumber_data_siap = True
      else:
        st.error("❌ Harap upload file arsip ZIP/RAR atau file Excel terlebih dahulu!")

      if sumber_data_siap:
        with st.spinner("Sedang memproses data omset & update kolom DATE..."):
          df_gabungan, logs = proses_gabung_data_dari_folder(
              temp_dir, jumlah_tanggal_t1, nama_bulan_t1
          )

          st.text_area(
              "📋 Log Status Proses Data:",
              value="\n".join(logs),
              height=150,
          )

          if df_gabungan is not None:
            try:
              template_path = os.path.join(
                  os.path.dirname(__file__), "template_omset.xlsx"
              )

              if not os.path.exists(template_path):
                st.error(
                    "❌ ERROR: File `template_omset.xlsx` belum di-upload ke"
                    " repository GitHub!"
                )
              else:
                output_file_path = os.path.join(
                    temp_dir, "hasil_template_omset.xlsx"
                )
                sh.copy(template_path, output_file_path)

                wb = openpyxl.load_workbook(output_file_path)
                ws = (
                    wb["ANALISA"]
                    if "ANALISA" in wb.sheetnames
                    else wb.active
                )

                eng_months = [
                    "",
                    "Jan",
                    "Feb",
                    "Mar",
                    "Apr",
                    "May",
                    "Jun",
                    "Jul",
                    "Aug",
                    "Sep",
                    "Oct",
                    "Nov",
                    "Dec",
                ]
                m_str = eng_months[angka_bulan_t1]

                _, total_hari_bulan_ini = calendar.monthrange(
                    pilihan_tahun_t1, angka_bulan_t1
                )

                start_row = 8
                col_index = 3

                for d in range(1, total_hari_bulan_ini + 1):
                  date_str = f"{d:02d}-{m_str}-{pilihan_tahun_t1}"
                  ws.cell(row=start_row + d - 1, column=col_index, value=date_str)

                wb.save(output_file_path)

                with pd.ExcelWriter(
                    output_file_path,
                    engine="openpyxl",
                    mode="a",
                    if_sheet_exists="replace",
                ) as writer:
                  df_gabungan.to_excel(
                      writer, sheet_name="statement", index=False
                  )

                st.success(
                    "🎉 Proses Selesai! Kolom DATE dari C8 sampai tanggal"
                    f" {total_hari_bulan_ini} berhasil di-generate penuh"
                    f" untuk bulan {pilihan_bulan_nama_t1} {pilihan_tahun_t1}."
                )

                with open(output_file_path, "rb") as f:
                  st.download_button(
                      label="📥 Download Hasil Template Excel",
                      data=f,
                      file_name=f"rekap_omset_{pilihan_bulan_nama_t1}_{pilihan_tahun_t1}.xlsx",
                      mime=(
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                      ),
                  )
            except Exception as e:
              st.error(f"❌ Terjadi kesalahan saat memproses Excel: {e}")


# ================= TAB 2: UPDATE EXISTING =================
with tab2:
  st.subheader("Update File Excel Target Tertentu via Web")
  st.write(
      "**Langkah 1**: Upload file Excel target beserta data harian sumbernya"
      " di bawah ini."
  )

  target_file_upload = st.file_uploader(
      "Upload File Excel Target (.xlsx/.xls):",
      type=["xlsx", "xls"],
      key="target_excel",
  )

  metode_upload_tab2 = st.radio(
      "Pilih Metode Upload Data Harian (Tab 2):",
      ["📁 Upload File Arsip (ZIP / RAR)", "📄 Upload Banyak File Excel (.xlsx/.xls)"],
      key="metode_tab2",
  )

  uploaded_archive_tab2 = None
  uploaded_files_tab2 = None

  if metode_upload_tab2 == "📁 Upload File Arsip (ZIP / RAR)":
    uploaded_archive_tab2 = st.file_uploader(
        "Upload file .zip atau .rar data harian untuk update target:",
        type=["zip", "rar"],
        key="archive_tab2",
    )
  else:
    uploaded_files_tab2 = st.file_uploader(
        "Upload file-file Excel harian untuk update target:",
        type=["xlsx", "xls"],
        accept_multiple_files=True,
        key="files_tab2",
    )

  st.markdown("")

  # --- PARAMETER PENGATURAN (DIBAWAH UPLOAD TAB 2) ---
  with st.container(border=True):
    st.markdown(
        "<h6 style='text-align: center; margin-bottom: 5px;'>⚙️ Parameter"
        " Pengaturan Data (Target)</h6>",
        unsafe_allow_html=True,
    )
    tp1, tp2, tp3 = st.columns(3)

    with tp1:
      jumlah_tanggal_t2 = st.selectbox(
          "Jumlah Tanggal (Proses):",
          options=list(range(1, 32)),
          index=0,
          key="jml_t2",
      )

    with tp2:
      pilihan_bulan_nama_t2 = st.selectbox(
          "Bulan:", options=list(bulan_dict.keys()), key="bln_t2"
      )
      nama_bulan_t2, angka_bulan_t2 = bulan_dict[pilihan_bulan_nama_t2]

    with tp3:
      pilihan_tahun_t2 = st.selectbox(
          "Tahun:",
          options=[2026, 2027, 2028, 2029, 2030],
          index=0,
          key="thn_t2",
      )

  st.markdown("")
  if st.button("🔄 Jalankan Update Excel Target", type="primary"):
    if not target_file_upload:
      st.error("❌ Harap upload file Excel target terlebih dahulu!")
    else:
      with tempfile.TemporaryDirectory() as temp_dir:
        sumber_data_siap = False

        if (
            metode_upload_tab2 == "📁 Upload File Arsip (ZIP / RAR)"
            and uploaded_archive_tab2 is not None
        ):
          with st.spinner("Mengekstrak file arsip..."):
            archive_path = os.path.join(temp_dir, uploaded_archive_tab2.name)
            with open(archive_path, "wb") as f:
              f.write(uploaded_archive_tab2.getbuffer())

            if uploaded_archive_tab2.name.endswith(".zip"):
              with zipfile.ZipFile(archive_path, "r") as zip_ref:
                zip_ref.extractall(temp_dir)
              sumber_data_siap = True
            elif uploaded_archive_tab2.name.endswith(".rar"):
              try:
                with rarfile.RarFile(archive_path, "r") as rar_ref:
                  rar_ref.extractall(temp_dir)
                sumber_data_siap = True
              except Exception as e:
                st.error(
                    f"❌ Gagal mengekstrak file RAR. Pastikan sistem mendukung unrar: {e}"
                )

        elif (
            metode_upload_tab2 == "📄 Upload Banyak File Excel (.xlsx/.xls)"
            and uploaded_files_tab2
        ):
          with st.spinner("Menyiapkan file yang di-upload..."):
            for uf in uploaded_files_tab2:
              temp_file_path = os.path.join(temp_dir, uf.name)
              with open(temp_file_path, "wb") as f:
                f.write(uf.getbuffer())
            sumber_data_siap = True
        else:
          st.error(
              "❌ Harap upload file arsip ZIP/RAR atau file Excel sumber harian"
              " terlebih dahulu!"
          )

        if sumber_data_siap:
          with st.spinner("Sedang memperbarui file Excel target..."):
            df_gabungan, logs = proses_gabung_data_dari_folder(
                temp_dir, jumlah_tanggal_t2, nama_bulan_t2
            )

            st.text_area(
                "📋 Log Status Update Target:",
                value="\n".join(logs),
                height=150,
            )

            if df_gabungan is not None:
              try:
                target_path_local = os.path.join(
                    temp_dir, target_file_upload.name
                )
                with open(target_path_local, "wb") as f:
                  f.write(target_file_upload.getbuffer())

                wb = openpyxl.load_workbook(target_path_local)
                ws = (
                    wb["ANALISA"]
                    if "ANALISA" in wb.sheetnames
                    else wb.active
                )

                eng_months = [
                    "",
                    "Jan",
                    "Feb",
                    "Mar",
                    "Apr",
                    "May",
                    "Jun",
                    "Jul",
                    "Aug",
                    "Sep",
                    "Oct",
                    "Nov",
                    "Dec",
                ]
                m_str = eng_months[angka_bulan_t2]

                _, total_hari_bulan_ini = calendar.monthrange(
                    pilihan_tahun_t2, angka_bulan_t2
                )
                start_row = 8
                col_index = 3

                for d in range(1, total_hari_bulan_ini + 1):
                  date_str = f"{d:02d}-{m_str}-{pilihan_tahun_t2}"
                  ws.cell(row=start_row + d - 1, column=col_index, value=date_str)

                wb.save(target_path_local)

                with pd.ExcelWriter(
                    target_path_local,
                    engine="openpyxl",
                    mode="a",
                    if_sheet_exists="replace",
                ) as writer:
                  df_gabungan.to_excel(
                      writer, sheet_name="statement", index=False
                  )

                st.success("🎉 Berhasil memperbarui file Excel target!")

                with open(target_path_local, "rb") as f:
                  st.download_button(
                      label="📥 Download File Target yang Telah Di-update",
                      data=f,
                      file_name=target_file_upload.name,
                      mime=(
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                      ),
                  )
              except Exception as e:
                st.error(
                    f"❌ Terjadi kesalahan saat memperbarui file target: {e}"
                )